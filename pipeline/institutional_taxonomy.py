"""
Institutional AI Tagging & Training Taxonomy loader and utilities.

Parses the structured institutional taxonomy Excel sheet containing:
- AI_Tagging_Taxonomy: 489 curated keywords with rich metadata (cues, expansions, model roles, label types)
- YOLO_Class_Index: 113 candidate visual/object classes for detection & segmentation
"""

import os
import re
import logging
from dataclasses import dataclass, field
from typing import List, Dict, Optional, Any, Set
import openpyxl

logger = logging.getLogger(__name__)


def _clean_str(val: Any) -> str:
    """Clean and strip a cell value, returning empty string if None."""
    if val is None:
        return ""
    return str(val).strip()


def _split_semicolons(val: Any) -> List[str]:
    """Split string by semicolons and return non-empty stripped items."""
    s = _clean_str(val)
    if not s:
        return []
    items = [x.strip() for x in s.split(';') if x.strip()]
    return items


@dataclass
class TaxonomyEntry:
    """Represents a single keyword entry from AI_Tagging_Taxonomy."""
    keyword_id: int
    source_sheet: str
    source_row: int
    project: str
    original_keyword: str
    normalized_keyword: str
    hard_drive: Optional[str]
    primary_category: str
    secondary_categories: List[str]
    keyword_type: str
    candidate_visual_class: Optional[str]
    visual_cues: List[str]
    context_associations: List[str]
    search_expansions: List[str]
    recommended_model_role: str
    evidence_basis: str
    training_use: str
    label_type: str  # 'object', 'context', 'scene', 'metadata', 'entity_metadata'
    priority: str
    inference_note: Optional[str] = None

    def get_embedding_prompts(self) -> List[str]:
        """
        Generate candidate text prompts for CLIP text encoding.
        Incorporates normalized keyword, visual cues, and expansions.
        """
        prompts: Set[str] = set()
        base = self.normalized_keyword.lower()
        if base:
            prompts.add(base)
            prompts.add(f"a photo of {base}")

        if self.candidate_visual_class:
            vc = self.candidate_visual_class.lower()
            prompts.add(vc)
            prompts.add(f"a photo of {vc}")

        for cue in self.visual_cues[:4]:
            c = cue.strip().lower()
            if c and len(c) > 2:
                prompts.add(f"a photo showing {c}")

        for exp in self.search_expansions[:4]:
            e = exp.strip().lower()
            if e and len(e) > 2:
                prompts.add(e)

        return list(prompts)


@dataclass
class YOLOClassEntry:
    """Represents a candidate visual/YOLO class from YOLO_Class_Index."""
    class_id: int
    canonical_class: str
    source_keyword_count: int
    source_keywords: List[str]
    primary_categories: List[str]
    projects: List[str]
    label_type: str
    recommended_models: str
    annotation_guidance: str
    status: str


class InstitutionalTaxonomy:
    """
    Manages loading and querying the institutional taxonomy.
    """
    def __init__(
        self,
        entries: List[TaxonomyEntry],
        yolo_classes: Optional[List[YOLOClassEntry]] = None
    ):
        self.entries = entries
        self.yolo_classes = yolo_classes or []
        self._by_id: Dict[int, TaxonomyEntry] = {e.keyword_id: e for e in entries}
        self._by_label_type: Dict[str, List[TaxonomyEntry]] = {}
        self._by_category: Dict[str, List[TaxonomyEntry]] = {}
        self._by_project: Dict[str, List[TaxonomyEntry]] = {}

        for entry in self.entries:
            self._by_label_type.setdefault(entry.label_type, []).append(entry)
            self._by_category.setdefault(entry.primary_category, []).append(entry)
            self._by_project.setdefault(entry.project, []).append(entry)

    @classmethod
    def from_excel(
        cls,
        excel_path: str,
        main_sheet: str = "AI_Tagging_Taxonomy",
        yolo_sheet: str = "YOLO_Class_Index"
    ) -> "InstitutionalTaxonomy":
        """Load and parse taxonomy from Excel workbook."""
        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"Taxonomy Excel file not found: {excel_path}")

        logger.info(f"Loading institutional taxonomy from {excel_path}...")
        wb = openpyxl.load_workbook(excel_path, data_only=True)

        entries: List[TaxonomyEntry] = []
        if main_sheet in wb.sheetnames:
            ws = wb[main_sheet]
            headers = [_clean_str(cell.value) for cell in ws[1]]
            header_map = {h: idx for idx, h in enumerate(headers) if h}

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row):
                    continue

                orig_kw = _clean_str(row[header_map.get("Original Keyword", 3)])
                if not orig_kw:
                    continue

                entry = TaxonomyEntry(
                    keyword_id=row_idx - 1,
                    source_sheet=_clean_str(row[header_map.get("Source Sheet", 0)]),
                    source_row=int(row[header_map.get("Source Row", 1)] or row_idx),
                    project=_clean_str(row[header_map.get("Project / Event", 2)]),
                    original_keyword=orig_kw,
                    normalized_keyword=_clean_str(row[header_map.get("Normalized Keyword", 4)]) or orig_kw,
                    hard_drive=_clean_str(row[header_map.get("Hard Drive", 5)]) or None,
                    primary_category=_clean_str(row[header_map.get("Primary Category", 8)]),
                    secondary_categories=_split_semicolons(row[header_map.get("Secondary Categories", 9)]),
                    keyword_type=_clean_str(row[header_map.get("Keyword Type", 10)]),
                    candidate_visual_class=_clean_str(row[header_map.get("Candidate Visual / YOLO Class", 11)]) or None,
                    visual_cues=_split_semicolons(row[header_map.get("Visual Cues / What May Appear", 12)]),
                    context_associations=_split_semicolons(row[header_map.get("Context Associations", 13)]),
                    search_expansions=_split_semicolons(row[header_map.get("Semantic Search Expansions", 14)]),
                    recommended_model_role=_clean_str(row[header_map.get("Recommended Model Role", 15)]),
                    evidence_basis=_clean_str(row[header_map.get("Evidence Basis", 16)]),
                    training_use=_clean_str(row[header_map.get("Training Use", 17)]),
                    label_type=_clean_str(row[header_map.get("Label Type", 18)]) or "context",
                    priority=_clean_str(row[header_map.get("Priority", 19)]) or "High",
                    inference_note=_clean_str(row[header_map.get("Inference / Evidence Note", 20)]) or None,
                )
                entries.append(entry)
            logger.info(f"Loaded {len(entries)} taxonomy entries from {main_sheet}.")
        else:
            logger.warning(f"Sheet {main_sheet} not found in {excel_path}.")

        yolo_classes: List[YOLOClassEntry] = []
        if yolo_sheet in wb.sheetnames:
            ws_yolo = wb[yolo_sheet]
            headers_y = [_clean_str(cell.value) for cell in ws_yolo[1]]
            hmap_y = {h: idx for idx, h in enumerate(headers_y) if h}

            for row_idx, row in enumerate(ws_yolo.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row):
                    continue
                cid_val = row[hmap_y.get("Class ID", 0)]
                if cid_val is None:
                    continue
                try:
                    cid = int(cid_val)
                except ValueError:
                    continue

                canonical = _clean_str(row[hmap_y.get("Canonical Candidate Visual Class", 1)])
                yc = YOLOClassEntry(
                    class_id=cid,
                    canonical_class=canonical,
                    source_keyword_count=int(row[hmap_y.get("Source Keyword Count", 2)] or 1),
                    source_keywords=_split_semicolons(row[hmap_y.get("Source Keywords", 3)]),
                    primary_categories=_split_semicolons(row[hmap_y.get("Primary Categories", 4)]),
                    projects=_split_semicolons(row[hmap_y.get("Projects / Events", 5)]),
                    label_type=_clean_str(row[hmap_y.get("Label Types", 7)]) or "object",
                    recommended_models=_clean_str(row[hmap_y.get("Recommended Models", 8)]),
                    annotation_guidance=_clean_str(row[hmap_y.get("Annotation Guidance", 9)]),
                    status=_clean_str(row[hmap_y.get("Status", 10)]),
                )
                yolo_classes.append(yc)
            logger.info(f"Loaded {len(yolo_classes)} YOLO classes from {yolo_sheet}.")

        return cls(entries=entries, yolo_classes=yolo_classes)

    def get_by_label_type(self, label_type: str) -> List[TaxonomyEntry]:
        """Filter entries by label type: 'object', 'context', 'scene', 'metadata', 'entity_metadata'."""
        return self._by_label_type.get(label_type, [])

    def get_by_category(self, category: str) -> List[TaxonomyEntry]:
        """Filter entries by primary category."""
        return self._by_category.get(category, [])

    def get_by_project(self, project: str) -> List[TaxonomyEntry]:
        """Filter entries by project/event name."""
        return self._by_project.get(project, [])

    def get_all_categories(self) -> List[str]:
        """Return sorted list of all unique primary categories."""
        return sorted([c for c in self._by_category.keys() if c])

    def get_all_projects(self) -> List[str]:
        """Return sorted list of all unique projects/events."""
        return sorted([p for p in self._by_project.keys() if p])

    def get_entry(self, keyword_id: int) -> Optional[TaxonomyEntry]:
        """Get entry by keyword ID."""
        return self._by_id.get(keyword_id)

    def __len__(self) -> int:
        return len(self.entries)
